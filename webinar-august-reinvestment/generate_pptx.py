#!/usr/bin/env python3
"""Generate charts and a PowerPoint deck for the August 2026 reinvestment webinar."""

from pathlib import Path

import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import nsmap, qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent
ASSETS = OUT / "assets"
ASSETS.mkdir(exist_ok=True)

NAVY = RGBColor(0x0A, 0x16, 0x28)
NAVY2 = RGBColor(0x12, 0x2B, 0x4A)
GOLD = RGBColor(0xD4, 0xA0, 0x17)
CREAM = RGBColor(0xF4, 0xEF, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x8A, 0x96, 0xA8)
TEAL = RGBColor(0x1A, 0x7A, 0x6D)
CORAL = RGBColor(0xC4, 0x5C, 0x4A)
SLATE = RGBColor(0x3D, 0x4F, 0x66)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, text, size, color, bold=False, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def add_text(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run() if p.runs else p.add_run(), text, size, color, bold, font)
    # python-pptx: first paragraph already has no run until we add one
    return box


def textbox(slide, l, t, w, h):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, space_after=6, font="Calibri"):
    p = tf.paragraphs[0] if not tf.paragraphs[0].text and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    set_run(run, text, size, color, bold, font)
    return p


def rect(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def round_rect(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def gold_bar(slide, l=0.7, t=0.55, w=0.55, h=0.08):
    return rect(slide, l, t, w, h, GOLD)


def footer(slide, n, total=16):
    tf = textbox(slide, 0.7, 7.12, 9, 0.28)
    para(tf, "August 2026  ·  Reinvesting Bond Maturities & Coupons  ·  For discussion with investors", 10, MUTED, space_after=0)
    tf2 = textbox(slide, 11.6, 7.12, 1.2, 0.28)
    para(tf2, f"{n}  /  {total}", 10, MUTED, align=PP_ALIGN.RIGHT, space_after=0)


def blank_dark():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def add_blank(prs):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    rect(slide, 0, 0, 13.333, 7.5, NAVY)
    return slide


def add_content(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    rect(slide, 0, 0, 13.333, 7.5, CREAM)
    rect(slide, 0, 0, 13.333, 0.08, NAVY)
    rect(slide, 0, 7.42, 13.333, 0.08, NAVY)
    return slide


def fmt_kes(n):
    if n >= 1_000_000:
        return f"KES {n/1_000_000:.2f}m"
    return f"KES {n:,.0f}"


def build_charts():
    years = list(range(0, 16))
    p = 1_000_000
    spend = [p if y == 0 else 0 for y in years]
    bank = [p * (1.02**y) for y in years]
    mmf = [p * (1.09**y) for y in years]
    fif = [p * (1.12**y) for y in years]

    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.facecolor": "#F4EFE4",
        "figure.facecolor": "#F4EFE4",
        "axes.labelcolor": "#0A1628",
        "xtick.color": "#3D4F66",
        "ytick.color": "#3D4F66",
        "text.color": "#0A1628",
    })

    fig, ax = plt.subplots(figsize=(12.2, 5.4), dpi=160)
    ax.plot(years, [s / 1e6 for s in spend], color="#C45C4A", lw=2.4, ls="--", label="Spent in year 1")
    ax.plot(years, [s / 1e6 for s in bank], color="#8A96A8", lw=2.6, label="Left in a 2% account")
    ax.plot(years, [s / 1e6 for s in mmf], color="#1A7A6D", lw=3.0, label="Money Market Fund @ 9%")
    ax.plot(years, [s / 1e6 for s in fif], color="#D4A017", lw=3.0, label="Fixed Income Fund @ 12%")
    ax.fill_between(years, [s / 1e6 for s in mmf], [s / 1e6 for s in bank], color="#1A7A6D", alpha=0.10)
    ax.set_xlim(0, 15)
    ax.set_ylim(0, 6.2)
    ax.set_xlabel("Years after August 2026", fontsize=11)
    ax.set_ylabel("Value of KES 1,000,000 (millions)", fontsize=11)
    ax.yaxis.set_major_formatter(mticker.FormatStrFormatter("%.1f"))
    ax.set_xticks(range(0, 16, 1))
    ax.grid(axis="y", color="#D9D2C5", lw=0.8)
    ax.legend(frameon=False, loc="upper left", fontsize=10)
    ax.annotate("KES 3.11m", xy=(10, fif[10] / 1e6), xytext=(8.2, 4.55),
                fontsize=9, color="#0A1628",
                arrowprops=dict(arrowstyle="->", color="#D4A017"))
    ax.annotate("KES 2.37m", xy=(10, mmf[10] / 1e6), xytext=(10.4, 1.85),
                fontsize=9, color="#0A1628",
                arrowprops=dict(arrowstyle="->", color="#1A7A6D"))
    fig.tight_layout()
    compounding_path = ASSETS / "compounding.png"
    fig.savefig(compounding_path, bbox_inches="tight")
    plt.close()

    # Bar comparison at 10 years
    fig, ax = plt.subplots(figsize=(10.6, 4.6), dpi=160)
    labels = ["Spent", "2% account", "Money Market\n9%", "Fixed Income\n12%"]
    values = [0, bank[10] / 1e6, mmf[10] / 1e6, fif[10] / 1e6]
    colors = ["#C45C4A", "#8A96A8", "#1A7A6D", "#D4A017"]
    bars = ax.bar(labels, values, color=colors, width=0.62, zorder=3)
    ax.set_ylabel("Value after 10 years (KES millions)", fontsize=11)
    ax.set_ylim(0, 3.8)
    ax.grid(axis="y", color="#D9D2C5", lw=0.8, zorder=0)
    for bar, val in zip(bars, [0, bank[10], mmf[10], fif[10]]):
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.08,
                "KES 0" if val == 0 else f"KES {val/1e6:.2f}m",
                ha="center", va="bottom", fontsize=10, fontweight="bold", color="#0A1628")
    fig.tight_layout()
    bars_path = ASSETS / "ten_year_bars.png"
    fig.savefig(bars_path, bbox_inches="tight")
    plt.close()

    return compounding_path, bars_path, {
        "bank5": bank[5], "mmf5": mmf[5], "fif5": fif[5],
        "bank10": bank[10], "mmf10": mmf[10], "fif10": fif[10],
        "bank15": bank[15], "mmf15": mmf[15], "fif15": fif[15],
    }


def build_deck(compounding_path, bars_path, nums):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    total = 16

    # 1 Title
    s = add_blank(prs)
    rect(s, 0, 0, 0.18, 7.5, GOLD)
    tf = textbox(s, 0.9, 1.55, 11.5, 0.4)
    para(tf, "INVESTOR WEBINAR  ·  WEDNESDAY 19 AUGUST 2026", 14, GOLD, bold=True, space_after=0)
    tf = textbox(s, 0.9, 2.05, 12, 1.8)
    para(tf, "Don’t let August cash\nleave your portfolio.", 40, WHITE, bold=True, space_after=8, font="Georgia")
    tf = textbox(s, 0.9, 4.15, 11, 0.9)
    para(tf, "Kenya’s domestic debt redemptions this month are almost double July’s.\nPut maturities and coupons back to work — in a Money Market Fund or a Fixed Income Fund.", 18, CREAM, space_after=0)
    tf = textbox(s, 0.9, 6.55, 11, 0.35)
    para(tf, "Illustrative briefing for investors  ·  Not a solicitation for any named fund", 12, MUTED, space_after=0)

    # 2 Agenda
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 11, 0.5)
    para(tf, "What we will cover", 28, NAVY, bold=True, font="Georgia")
    items = [
        ("01", "August’s cash calendar", "Which bonds are redeeming, which coupons land, and why the month is unusually large."),
        ("02", "The cost of spending it", "A compounding picture of KES 1,000,000 spent vs reinvested over 5, 10 and 15 years."),
        ("03", "Two natural homes for the cash", "When a Money Market Fund is the right parking bay — and when a Fixed Income Fund should do the work."),
        ("04", "What to do this week", "A simple decision path before proceeds hit the linked bank account and disappear into spending."),
    ]
    for i, (num, title, body) in enumerate(items):
        top = 1.45 + i * 1.25
        round_rect(s, 0.7, top, 11.9, 1.12, WHITE)
        tf = textbox(s, 0.95, top + 0.22, 0.8, 0.6)
        para(tf, num, 22, GOLD, bold=True, space_after=0)
        tf = textbox(s, 1.9, top + 0.16, 10.3, 0.4)
        para(tf, title, 18, NAVY, bold=True, space_after=0)
        tf = textbox(s, 1.9, top + 0.55, 10.3, 0.45)
        para(tf, body, 14, SLATE, space_after=0)
    footer(s, 2, total)

    # 3 Why August
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.5)
    para(tf, "August is a cash-flow month, not a quiet month", 26, NAVY, bold=True, font="Georgia")
    cards = [
        ("KES 294bn", "Domestic debt maturing in August 2026", "Almost double July’s KES 152bn of T-bill and T-bond redemptions. (SIB, citing CBK)"),
        ("17 August", "Two government papers redeem", "FXD1/2016/010 matures in full. IFB1/2020/011 sits on CBK’s 17 Aug redemption calendar (scheduled amortisation)."),
        ("Coupons + bills", "Interest and T-bills land all month", "Feb/August coupon bonds pay this month. Weekly T-bills redeem on 17, 24 and 31 August."),
    ]
    for i, (stat, title, body) in enumerate(cards):
        left = 0.7 + i * 4.15
        round_rect(s, left, 1.5, 3.95, 4.55, WHITE)
        rect(s, left, 1.5, 3.95, 0.08, GOLD)
        tf = textbox(s, left + 0.25, 1.85, 3.45, 1.1)
        para(tf, stat, 26, NAVY, bold=True, font="Georgia")
        tf = textbox(s, left + 0.25, 3.05, 3.45, 0.8)
        para(tf, title, 16, TEAL, bold=True)
        tf = textbox(s, left + 0.25, 3.85, 3.45, 1.8)
        para(tf, body, 13, SLATE)
    footer(s, 3, total)

    # 4 Bonds maturing
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.45)
    para(tf, "Bonds on the August redemption calendar", 26, NAVY, bold=True, font="Georgia")
    tf = textbox(s, 0.7, 1.2, 12, 0.4)
    para(tf, "Source: CBK redemption calendar and original prospectuses. Outstanding amounts vary after switches and reopenings.", 13, SLATE)

    headers = ["Issue", "Type", "What happens", "Coupon", "Tax"]
    rows = [
        ["FXD1/2016/010", "10-year Treasury bond", "Matures in full — 17 Aug 2026", "15.039%", "10% WHT"],
        ["IFB1/2020/011", "11-year Infrastructure bond", "On CBK’s 17 Aug redemption list (typical 50% year-6 amortisation; original final maturity 2031)", "10.90%", "Tax-exempt"],
        ["Weekly T-bills", "91 / 182 / 364-day", "Redemptions 17, 24 and 31 August 2026", "Discount", "15% WHT"],
        ["FXD1/2012/015 + 3 T-bills", "Switch offer", "Optional switch into FXD4/2019/010 — bids by 24 Aug, settles 26 Aug", "11.00% source", "10% WHT"],
    ]
    col_w = [2.2, 2.3, 4.4, 1.5, 1.4]
    x0, y0 = 0.7, 1.7
    # header
    rect(s, x0, y0, 11.9, 0.48, NAVY)
    x = x0
    for h, w in zip(headers, col_w):
        tf = textbox(s, x + 0.08, y0 + 0.08, w - 0.1, 0.35)
        para(tf, h, 12, WHITE, bold=True, space_after=0)
        x += w
    for r, row in enumerate(rows):
        y = y0 + 0.48 + r * 0.95
        fill = WHITE if r % 2 == 0 else RGBColor(0xEA, 0xE4, 0xD6)
        rect(s, x0, y, 11.9, 0.95, fill)
        x = x0
        for j, (cell, w) in enumerate(zip(row, col_w)):
            tf = textbox(s, x + 0.08, y + 0.18, w - 0.12, 0.65)
            para(tf, cell, 12, NAVY, bold=(j == 0), space_after=0)
            x += w
    footer(s, 4, total)

    # 5 Coupons
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.45)
    para(tf, "Coupons landing in August — even if your bond is not maturing", 24, NAVY, bold=True, font="Georgia")
    tf = textbox(s, 0.7, 1.25, 12, 0.55)
    para(tf, "Kenya Treasury bonds pay interest twice a year. Papers issued on a February / August cycle credit coupons this month. That cash is easy to treat as “extra” and spend. It is still capital that can compound.", 15, SLATE)

    examples = [
        ("10 August", "FXD1/2021/020", "13.444% coupon  ·  10% WHT  ·  matures July 2041"),
        ("17 August", "FXD1/2016/010", "Final coupon + 100% principal  ·  15.039%"),
        ("Feb / Aug cycle", "Several FXD and IFB issues", "Holders receive semi-annual interest into DhowCSD / the linked bank"),
    ]
    for i, (when, name, detail) in enumerate(examples):
        top = 2.05 + i * 1.25
        round_rect(s, 0.7, top, 11.9, 1.12, WHITE)
        tf = textbox(s, 0.95, top + 0.22, 2.4, 0.7)
        para(tf, when, 16, GOLD, bold=True)
        tf = textbox(s, 3.5, top + 0.18, 8.7, 0.4)
        para(tf, name, 18, NAVY, bold=True, space_after=0)
        tf = textbox(s, 3.5, top + 0.58, 8.7, 0.4)
        para(tf, detail, 14, SLATE, space_after=0)
    footer(s, 5, total)

    # 6 The leak
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.45)
    para(tf, "What usually happens to the money", 26, NAVY, bold=True, font="Georgia")
    steps = [
        ("1", "Credit", "CBK pays principal and coupons into DhowCSD. Cash is then available to the linked bank account."),
        ("2", "The quiet leak", "Once it sits in a transactional account it looks like income. School fees, a trip, a new car — all reasonable, all terminal for compounding."),
        ("3", "Inflation", "At 6.5% inflation (July 2026), cash that earns ~0–2% loses purchasing power every month it waits."),
        ("4", "The better default", "Treat every redemption as a reinvestment event first. Spend only what you planned to spend — from a budget, not from a maturity."),
    ]
    for i, (n, title, body) in enumerate(steps):
        left = 0.7 + (i % 2) * 6.2
        top = 1.45 + (i // 2) * 2.45
        round_rect(s, left, top, 5.9, 2.25, WHITE)
        tf = textbox(s, left + 0.3, top + 0.3, 0.7, 0.5)
        para(tf, n, 22, GOLD, bold=True, space_after=0)
        tf = textbox(s, left + 1.0, top + 0.35, 4.5, 0.4)
        para(tf, title, 18, NAVY, bold=True, space_after=0)
        tf = textbox(s, left + 0.3, top + 0.95, 5.3, 1.05)
        para(tf, body, 14, SLATE)
    footer(s, 6, total)

    # 7 Compounding hero
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.62, 12, 0.4)
    para(tf, "KES 1,000,000 received this August — four paths", 24, NAVY, bold=True, font="Georgia")
    tf = textbox(s, 0.7, 1.05, 12, 0.35)
    para(tf, "Illustrative only. 9% and 12% are round numbers near recent MMF / fixed-income industry averages, not a promise.", 12, SLATE)
    s.shapes.add_picture(str(compounding_path), Inches(0.55), Inches(1.4), Inches(12.2), Inches(5.4))
    footer(s, 7, total)

    # 8 10 year bars + table
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.62, 12, 0.4)
    para(tf, "Same shilling. Three futures. One empty pocket.", 24, NAVY, bold=True, font="Georgia")
    s.shapes.add_picture(str(bars_path), Inches(0.5), Inches(1.15), Inches(7.4), Inches(3.4))

    # table of 5/10/15
    headers = ["Horizon", "2% account", "MMF 9%", "FIF 12%", "Gap vs spending"]
    data = [
        ["5 years", f"{nums['bank5']/1e6:.2f}m", f"{nums['mmf5']/1e6:.2f}m", f"{nums['fif5']/1e6:.2f}m", f"{nums['fif5']/1e6:.2f}m"],
        ["10 years", f"{nums['bank10']/1e6:.2f}m", f"{nums['mmf10']/1e6:.2f}m", f"{nums['fif10']/1e6:.2f}m", f"{nums['fif10']/1e6:.2f}m"],
        ["15 years", f"{nums['bank15']/1e6:.2f}m", f"{nums['mmf15']/1e6:.2f}m", f"{nums['fif15']/1e6:.2f}m", f"{nums['fif15']/1e6:.2f}m"],
    ]
    x0, y0 = 0.7, 4.65
    col_w = [2.0, 2.3, 2.3, 2.3, 3.0]
    rect(s, x0, y0, 11.9, 0.38, NAVY)
    x = x0
    for h, w in zip(headers, col_w):
        tf = textbox(s, x + 0.08, y0 + 0.05, w - 0.1, 0.28)
        para(tf, h, 11, WHITE, bold=True, space_after=0)
        x += w
    for r, row in enumerate(data):
        y = y0 + 0.38 + r * 0.48
        rect(s, x0, y, 11.9, 0.48, WHITE if r % 2 == 0 else RGBColor(0xEA, 0xE4, 0xD6))
        x = x0
        for cell, w in zip(row, col_w):
            tf = textbox(s, x + 0.08, y + 0.08, w - 0.1, 0.32)
            para(tf, cell, 13, NAVY, bold=True, space_after=0)
            x += w
    footer(s, 8, total)

    # 9 Why reinvest
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.45)
    para(tf, "Why reinvest instead of spend", 26, NAVY, bold=True, font="Georgia")
    points = [
        ("Maturity is not a bonus", "The principal was already yours. Spending it shrinks the machine that was paying you 10–15% coupons."),
        ("Coupons are easy to hide", "A few hundred thousand shillings of interest does not feel like a decision. Repeated twice a year, it is the whole return."),
        ("Cash is taxed by inflation", "6.5% inflation vs a 2% account is a 4.5% real loss. A ~9% MMF still compounds in real terms."),
        ("You cannot recapture skipped years", "The steep part of the compounding curve is later. Ten idle years cannot be ‘caught up’ with one good year."),
        ("Liquidity is already built in", "You do not need to spend the bond to stay liquid. An MMF is typically T+2 to T+4 with a KES 1,000 minimum."),
        ("Replace the expired contract", "FXD1/2016/010 paid 15.039%. That contract ends on 17 August. A fund puts you back into a diversified book of bills and bonds the next day."),
    ]
    for i, (title, body) in enumerate(points):
        left = 0.7 + (i % 3) * 4.15
        top = 1.4 + (i // 3) * 2.55
        round_rect(s, left, top, 3.95, 2.35, WHITE)
        tf = textbox(s, left + 0.22, top + 0.25, 3.5, 0.7)
        para(tf, title, 15, NAVY, bold=True)
        tf = textbox(s, left + 0.22, top + 0.95, 3.5, 1.2)
        para(tf, body, 13, SLATE)
    footer(s, 9, total)

    # 10 Two funds overview
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.4)
    para(tf, "Two CMA-licensed homes for the proceeds", 26, NAVY, bold=True, font="Georgia")
    # two columns
    round_rect(s, 0.7, 1.4, 5.85, 5.35, WHITE)
    rect(s, 0.7, 1.4, 5.85, 0.1, TEAL)
    tf = textbox(s, 0.95, 1.7, 5.4, 0.4)
    para(tf, "Money Market Fund", 22, TEAL, bold=True, font="Georgia")
    tf = textbox(s, 0.95, 2.2, 5.4, 0.5)
    para(tf, "Park it. Stay liquid. Beat the current account.", 14, SLATE)
    mmf_bits = [
        "Holds T-bills, short bonds, call and fixed deposits",
        "Typical industry yield ~9% gross / ~7.6% net (H1 2026 avg.)",
        "Interest is calculated daily, paid after 15% WHT",
        "Access usually within 2–4 working days",
        "Best when the money may be needed inside 12 months",
        "KES 442bn of industry MMF AUM as of March 2026",
    ]
    tf = textbox(s, 0.95, 2.75, 5.4, 3.7)
    first = True
    for b in mmf_bits:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(10)
        run = p.add_run()
        set_run(run, "▸  " + b, 14, NAVY)

    round_rect(s, 6.8, 1.4, 5.85, 5.35, WHITE)
    rect(s, 6.8, 1.4, 5.85, 0.1, GOLD)
    tf = textbox(s, 7.05, 1.7, 5.4, 0.4)
    para(tf, "Fixed Income Fund", 22, GOLD, bold=True, font="Georgia")
    tf = textbox(s, 7.05, 2.2, 5.4, 0.5)
    para(tf, "Put duration back to work. Replace the expired bond.", 14, SLATE)
    fif_bits = [
        "Holds a ladder of Treasury and quality corporate bonds",
        "Category average ~11.9% gross / ~10.1% net (June 2026)",
        "More yield than an MMF; more price movement too",
        "Sensible holding period: 12 months and beyond",
        "~250 bps extra vs the MMF average in mid-2026",
        "A closer substitute for the bond that just matured",
    ]
    tf = textbox(s, 7.05, 2.75, 5.4, 3.7)
    first = True
    for b in fif_bits:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(10)
        run = p.add_run()
        set_run(run, "▸  " + b, 14, NAVY)
    footer(s, 10, total)

    # 11 Case MMF
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.4)
    para(tf, "The case for the Money Market Fund", 26, NAVY, bold=True, font="Georgia")
    tf = textbox(s, 0.7, 1.2, 12, 0.45)
    para(tf, "Use it when the August cash is a bridge — school fees in January, a property completion, or ‘I have not decided yet’.", 15, SLATE)
    reasons = [
        ("Beats idle cash immediately", "91-day T-bills are ~8.8%. A decent MMF often matches or beats that, with daily compounding and no auction paperwork."),
        ("No need to pick a bond", "You do not have to bid the 24 August switch or the next IFB. The manager already holds a book of bills and short paper."),
        ("You stay in control", "Units can usually be withdrawn in a few days. That is the point: the money is working, but it is not locked."),
        ("A disciplined default", "Standing instruction: every DhowCSD credit → MMF. Spending then requires a conscious withdrawal, not an accident."),
        ("Professional cash management", "Custodian + trustee + CMA licence. This is not a Sacco promise or a WhatsApp ‘investment’."),
        ("USD option exists", "If the proceeds are dollar-linked (or you want a USD sleeve), dollar MMFs returned ~4.3% net in H1 2026."),
    ]
    for i, (title, body) in enumerate(reasons):
        left = 0.7 + (i % 3) * 4.15
        top = 1.8 + (i // 3) * 2.4
        round_rect(s, left, top, 3.95, 2.2, WHITE)
        tf = textbox(s, left + 0.22, top + 0.22, 3.5, 0.55)
        para(tf, title, 15, TEAL, bold=True)
        tf = textbox(s, left + 0.22, top + 0.8, 3.5, 1.2)
        para(tf, body, 13, SLATE)
    footer(s, 11, total)

    # 12 Case FIF
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.4)
    para(tf, "The case for the Fixed Income Fund", 26, NAVY, bold=True, font="Georgia")
    tf = textbox(s, 0.7, 1.2, 12, 0.5)
    para(tf, "Use it when the August cash was long-term capital — the 10-year bond that just matured should not become a 91-day afterthought.", 15, SLATE)
    reasons = [
        ("Replace the expired coupon engine", "FXD1/2016/010 paid 15%. New 10-year FXD paper is around 12.8%. A bond fund keeps you in that market without timing one auction."),
        ("You get a ladder, not a single name", "One matured issue becomes a diversified book — several tenors, reopenings, and (in some funds) quality corporates."),
        ("The extra 2–3 points compound", "On KES 1m, 250 bps extra vs an MMF is ~KES 25,000 a year at the start — and more as the base grows."),
        ("Yields are still historically useful", "CBR is 8.75%. Long bonds still clear ~12–14%. Locking duration now is how you keep income if policy rates drift lower."),
        ("Tax-aware IFB exposure", "Infrastructure bonds pay coupons with no WHT. Funds that hold IFBs pass through some of that tax advantage."),
        ("Stay invested through the switch", "You do not have to decide by 24 August whether to switch FXD1/2012/015. The fund can hold duration for you."),
    ]
    for i, (title, body) in enumerate(reasons):
        left = 0.7 + (i % 3) * 4.15
        top = 1.85 + (i // 3) * 2.35
        round_rect(s, left, top, 3.95, 2.15, WHITE)
        tf = textbox(s, left + 0.22, top + 0.18, 3.5, 0.6)
        para(tf, title, 14, GOLD, bold=True)
        tf = textbox(s, left + 0.22, top + 0.78, 3.5, 1.2)
        para(tf, body, 12, SLATE)
    footer(s, 12, total)

    # 13 How to choose
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.4)
    para(tf, "How to choose — a 60-second test", 26, NAVY, bold=True, font="Georgia")
    headers = ["If this is true…", "Then use"]
    rows = [
        ["I may need this money within 12 months (fees, construction drawdown, emergency reserve).", "Money Market Fund"],
        ["I have not decided yet, and I refuse to leave it in the current account over the weekend.", "Money Market Fund  (default)"],
        ["This was a 5–10 year bond. I still want income, and I can leave it for a year or more.", "Fixed Income Fund"],
        ["I want both: spend nothing, keep 30–40% liquid, put the rest back into duration.", "Split: MMF + Fixed Income"],
        ["I am sitting on FXD1/2012/015 and the T-bills in the 26 August switch.", "Switch and/or Fixed Income Fund"],
        ["Dollar liabilities (school fees abroad, import invoice).", "USD Money Market sleeve"],
    ]
    x0, y0 = 0.7, 1.35
    rect(s, x0, y0, 11.9, 0.42, NAVY)
    tf = textbox(s, 0.85, y0 + 0.06, 8.2, 0.3)
    para(tf, "If this is true…", 13, WHITE, bold=True, space_after=0)
    tf = textbox(s, 9.2, y0 + 0.06, 3.2, 0.3)
    para(tf, "Then use", 13, WHITE, bold=True, space_after=0)
    for r, row in enumerate(rows):
        y = y0 + 0.42 + r * 0.78
        rect(s, x0, y, 11.9, 0.78, WHITE if r % 2 == 0 else RGBColor(0xEA, 0xE4, 0xD6))
        tf = textbox(s, 0.85, y + 0.12, 8.2, 0.55)
        para(tf, row[0], 13, NAVY, space_after=0)
        tf = textbox(s, 9.2, y + 0.18, 3.2, 0.45)
        para(tf, row[1], 13, TEAL, bold=True, space_after=0)
    footer(s, 13, total)

    # 14 This week
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.4)
    para(tf, "This week’s action list", 26, NAVY, bold=True, font="Georgia")
    actions = [
        ("Mon–Tue", "Open DhowCSD. Note every August credit — maturity, amortisation, coupon, T-bill."),
        ("Before 17 Aug", "If FXD1/2016/010 or IFB1/2020/011 is paying you, decide MMF vs FIF before the money hits the bank."),
        ("By 24 Aug, 10:00", "Holders of FXD1/2012/015 and the three source T-bills: accept or ignore the switch into FXD4/2019/010."),
        ("Same day as credit", "Move the amount you will not spend in 90 days. Standing order into the fund beats willpower."),
        ("Split if unsure", "Example: 40% MMF (near cash) + 60% Fixed Income (replace the bond). Review in 90 days."),
        ("Do not wait for ‘the next auction’", "Funds subscribe for you. Cash that waits for a perfect bond often becomes a school-fees withdrawal."),
    ]
    for i, (when, body) in enumerate(actions):
        top = 1.3 + i * 0.88
        round_rect(s, 0.7, top, 11.9, 0.8, WHITE)
        tf = textbox(s, 0.95, top + 0.22, 2.6, 0.4)
        para(tf, when, 14, GOLD, bold=True, space_after=0)
        tf = textbox(s, 3.7, top + 0.18, 8.6, 0.5)
        para(tf, body, 14, NAVY, space_after=0)
    footer(s, 14, total)

    # 15 Takeaways
    s = add_content(prs)
    gold_bar(s)
    tf = textbox(s, 0.7, 0.7, 12, 0.4)
    para(tf, "If you remember only five things", 26, NAVY, bold=True, font="Georgia")
    takes = [
        "August 2026 is a redemption month — KES 294bn of domestic paper is rolling off.",
        "FXD1/2016/010 matures 17 August. Coupons on the Feb/August cycle land too.",
        "Spending the principal ends the compounding. A 2% account is a slow version of the same mistake.",
        "Default: Money Market Fund if the horizon is short. Fixed Income Fund if you are replacing a long bond.",
        "Move the money the day it arrives. The decision is this week, not ‘after the holidays’.",
    ]
    for i, t in enumerate(takes):
        top = 1.4 + i * 0.95
        round_rect(s, 0.7, top, 11.9, 0.85, WHITE)
        tf = textbox(s, 0.95, top + 0.2, 0.7, 0.5)
        para(tf, f"{i+1:02d}", 20, GOLD, bold=True, space_after=0)
        tf = textbox(s, 1.85, top + 0.22, 10.4, 0.5)
        para(tf, t, 16, NAVY, space_after=0)
    footer(s, 15, total)

    # 16 Close
    s = add_blank(prs)
    rect(s, 0, 0, 0.18, 7.5, GOLD)
    tf = textbox(s, 0.9, 1.6, 11.5, 0.4)
    para(tf, "QUESTIONS  ·  NEXT STEPS", 14, GOLD, bold=True, space_after=0)
    tf = textbox(s, 0.9, 2.1, 12, 1.4)
    para(tf, "We can open the fund\nthe same week the bond pays.", 32, WHITE, bold=True, font="Georgia")
    tf = textbox(s, 0.9, 4.0, 11, 1.1)
    para(tf, "Bring your DhowCSD statement. We will map each August credit to either\na Money Market Fund, a Fixed Income Fund, or a planned withdrawal.", 18, CREAM)
    tf = textbox(s, 0.9, 5.5, 11.5, 1.2)
    para(tf, "Disclaimer: This briefing is educational. It is not a prospectus, not tax advice, and not a recommendation of any named unit trust. Yields cited are industry snapshots (CBK, CMA-licensed fund tables, SIB weekly) and will change. Past returns do not predict future returns. Invest only in CMA-licensed funds. Confirm fact sheets, fees, WHT and redemption timelines with the manager before you subscribe.", 12, MUTED)

    out = OUT / "August_2026_Reinvesting_Bond_Proceeds.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    compounding, bars, nums = build_charts()
    path = build_deck(compounding, bars, nums)
    print("Wrote", path)
    for k, v in nums.items():
        print(f"  {k}: {v:,.0f}")
